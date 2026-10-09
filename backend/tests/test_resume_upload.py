"""简历文件上传：PDF / docx / txt(md) 提取入库，原文件落盘，类型与空文件拒绝。"""

from __future__ import annotations

from io import BytesIO

from fastapi.testclient import TestClient

from app.config import get_settings


def _make_pdf(text: str) -> bytes:
    """手写一个合法最小单页 PDF（xref 偏移逐个计算，pypdf 可正常解析）。"""
    content = f"BT /F1 24 Tf 100 700 Td ({text}) Tj ET".encode("latin-1")
    objs = [
        b"<< /Type /Catalog /Pages 2 0 R >>",
        b"<< /Type /Pages /Kids [3 0 R] /Count 1 >>",
        b"<< /Type /Page /Parent 2 0 R /MediaBox [0 0 612 792] /Contents 4 0 R"
        b" /Resources << /Font << /F1 5 0 R >> >> >>",
        b"<< /Length %d >>\nstream\n%s\nendstream" % (len(content), content),
        b"<< /Type /Font /Subtype /Type1 /BaseFont /Helvetica >>",
    ]
    out = bytearray(b"%PDF-1.4\n")
    offsets = [0]
    for i, obj in enumerate(objs, start=1):
        offsets.append(len(out))
        out += b"%d 0 obj\n" % i
        out += obj
        out += b"\nendobj\n"
    xref_pos = len(out)
    out += b"xref\n0 %d\n" % (len(objs) + 1)
    out += b"0000000000 65535 f \n"
    for off in offsets[1:]:
        out += b"%010d 00000 n \n" % off
    out += b"trailer\n<< /Size %d /Root 1 0 R >>\nstartxref\n%d\n%%%%EOF" % (
        len(objs) + 1,
        xref_pos,
    )
    return bytes(out)


def _make_docx() -> bytes:
    from docx import Document

    document = Document()
    document.add_paragraph("视觉应用工程师简历")
    document.add_paragraph("负责相机选型与打光方案")
    table = document.add_table(rows=1, cols=2)
    table.rows[0].cells[0].text = "项目"
    table.rows[0].cells[1].text = "误检率 5% → 1%"
    buffer = BytesIO()
    document.save(buffer)
    return buffer.getvalue()


def _upload(client: TestClient, filename: str, payload: bytes, **data: str) -> object:
    return client.post(
        "/api/resumes/upload",
        files={"file": (filename, payload, "application/octet-stream")},
        data=data,
    )


def test_upload_txt_utf8(client: TestClient) -> None:
    body = "我叫王工，两年机器视觉经验。\n\n负责缺陷检测项目。".encode("utf-8")
    response = _upload(client, "王工的简历.txt", body)
    assert response.status_code == 201, response.text
    resume = response.json()
    assert resume["source"] == "upload"
    assert resume["file_name"] == "王工的简历.txt"
    assert "两年机器视觉经验" in resume["content_text"]
    assert "负责缺陷检测项目" in resume["content_text"]
    # 原文件落盘可追溯
    settings = get_settings()
    stored = list((settings.upload_dir / "resumes").glob(f"{resume['id']}_*"))
    assert len(stored) == 1


def test_upload_txt_gbk(client: TestClient) -> None:
    body = "张工\n视觉工程师\n主导九点标定".encode("gbk")
    response = _upload(client, "zhang.txt", body, name="张工简历")
    assert response.status_code == 201, response.text
    resume = response.json()
    assert resume["name"] == "张工简历"
    assert "视觉工程师" in resume["content_text"]
    assert "主导九点标定" in resume["content_text"]


def test_upload_docx_extracts_paragraphs_and_tables(client: TestClient) -> None:
    response = _upload(client, "resume.docx", _make_docx())
    assert response.status_code == 201, response.text
    resume = response.json()
    assert resume["source"] == "upload"
    assert "视觉应用工程师简历" in resume["content_text"]
    assert "误检率 5% → 1%" in resume["content_text"]


def test_upload_pdf_extracts_text(client: TestClient) -> None:
    response = _upload(client, "cv.pdf", _make_pdf("Machine Vision Engineer"))
    assert response.status_code == 201, response.text
    resume = response.json()
    assert "Machine Vision Engineer" in resume["content_text"]


def test_upload_rejects_unsupported_suffix(client: TestClient) -> None:
    response = _upload(client, "old.doc", b"\xd0\xcf\x11\xe0\xa1\xb1\x1a\xe1")
    assert response.status_code == 422
    assert "另存为" in response.json()["detail"]


def test_upload_rejects_empty_file(client: TestClient) -> None:
    response = _upload(client, "empty.txt", b"")
    assert response.status_code == 422


def test_upload_bumps_version_within_family(client: TestClient) -> None:
    first = _upload(client, "same.txt", "第一版".encode("utf-8")).json()
    second = _upload(client, "same.txt", "第二版".encode("utf-8")).json()
    assert first["version"] == 1
    assert second["version"] == 2
    assert second["name"] == "same"


def test_upload_unknown_target_job(client: TestClient) -> None:
    response = _upload(client, "a.txt", "内容".encode("utf-8"), target_job_id="999")
    assert response.status_code == 404


def test_upload_md_supported(client: TestClient) -> None:
    body = "# 简历\n\n- 项目 A：亚像素定位\n- 项目 B：模板匹配".encode("utf-8")
    response = _upload(client, "resume.md", body)
    assert response.status_code == 201, response.text
    resume = response.json()
    assert "亚像素定位" in resume["content_text"]
